"""Writing the output workbook with the template's own formatting.

pandas' `to_excel` writes bare cells, which loses the colours, fonts, borders,
column widths and freeze panes that make the template readable. Instead the
template file is *copied* to the output path - so everything not touched here
survives untouched - and only the data rows are replaced.

Formatting varies per column rather than per sheet (Tahoma 8/9/10pt headers over
different fills, the yellow Fingerprint column, 8pt vs 11pt data cells), so the
style is captured from the template's first data row column by column and
re-applied to every row written.
"""
from __future__ import annotations

import re
import shutil
from copy import copy
from pathlib import Path

from openpyxl import load_workbook

# Row 1 is the header; the template's first data row supplies the styling for
# every row written.
HEADER_ROW = 1
STYLE_ROW = 2

_INT_RE = re.compile(r"-?\d+")
_FLOAT_RE = re.compile(r"-?\d*\.\d+")


def coerce(value):
    """Text -> the type the template stores it as.

    The template holds eNBId, VLANs and ARFCNs as real numbers and only the
    identifiers as text. Writing everything as text would put a green
    'number stored as text' marker on most of the sheet and read differently to
    the reference file, so numeric-looking values are converted back.

    Leading zeros are kept as text: those are identifiers whose width matters,
    not quantities.
    """
    if value is None:
        return None
    s = str(value).strip()
    if not s:
        return None
    if _INT_RE.fullmatch(s):
        digits = s.lstrip("-")
        if len(digits) > 1 and digits.startswith("0"):
            return s
        try:
            return int(s)
        except ValueError:
            return s
    if _FLOAT_RE.fullmatch(s):
        try:
            return float(s)
        except ValueError:
            return s
    return s


def _column_styles(ws, ncols: int):
    """The style of each column, taken from the template's first data row."""
    if ws.max_row < STYLE_ROW:
        return {}
    return {c: copy(ws.cell(row=STYLE_ROW, column=c)._style)
            for c in range(1, ncols + 1)}


def write_workbook(template_path: Path, out_path: Path,
                   sheet_rows: dict[str, list[dict]],
                   sheet_columns: dict[str, list[str]]) -> Path:
    """Copy the template and replace its data rows with `sheet_rows`.

    `sheet_rows` maps sheet name -> list of dicts keyed by template column.
    """
    out_path.parent.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(template_path, out_path)

    wb = load_workbook(out_path)
    for sheet, rows in sheet_rows.items():
        ws = wb[sheet]
        columns = sheet_columns[sheet]
        styles = _column_styles(ws, len(columns))
        height = ws.row_dimensions[STYLE_ROW].height

        # Drop the template's sample data, keeping the header and everything
        # sheet-level (widths, freeze panes, tab colour).
        if ws.max_row >= STYLE_ROW:
            ws.delete_rows(STYLE_ROW, ws.max_row - HEADER_ROW)

        for i, row in enumerate(rows):
            excel_row = STYLE_ROW + i
            for ci, name in enumerate(columns, start=1):
                cell = ws.cell(row=excel_row, column=ci)
                cell.value = coerce(row.get(name))
                if ci in styles:
                    cell._style = copy(styles[ci])
            if height is not None:
                ws.row_dimensions[excel_row].height = height

    wb.save(out_path)
    return out_path
