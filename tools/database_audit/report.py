"""Excel report of the Database Audit - sheets, columns and styling of gpl_audit_tool_V2's report.

Long sheets (one row per parameter): LTE_GPL_AUDIT, NR_GPL_AUDIT, OTHER_MO_AUDIT.
Wide sheets (one row per MO, 'pre|post' where they differ, as V2): eNBInfo, gNBInfo, LTE_CELL_DATA, NR_CELL_DATA,
FREQ_AUDIT, LTE_EutranfreqRelation, NR_GutranfreqRelation, GeranFreqRelation, CellRelation - one titled table per
MO class (a full comparison mixes classes with different attributes on one sheet).
"""
from __future__ import annotations

import re
from collections import OrderedDict
from pathlib import Path

import pandas as pd
from openpyxl.utils import get_column_letter

from tools.gpl_audit.log_parser.output import (_TABLE_GAP_ROWS, _TITLE_ALIGN, _TITLE_FILL, _TITLE_FONT, _format_sheet,
                                                 _style_block)

from .audit import (CELL_REL, EUTRAN_FREQ, EUTRAN_FREQ_REL, FEATURE, GERAN_FREQ_REL, GUTRAN_FREQ_REL, LTE_GPL, NR_GPL,
                    CELL_CLASSES, OTHER, SSB_FREQ, AuditResult, classes_of, section_of, text)
from .corrections import MISSING, Plan, count_commands

WIDE_SHEETS = OrderedDict([
    ("FREQ_AUDIT", (EUTRAN_FREQ, SSB_FREQ)),
    ("LTE_EutranfreqRelation", (EUTRAN_FREQ_REL,)),
    ("NR_GutranfreqRelation", (GUTRAN_FREQ_REL,)),
    ("GeranFreqRelation", (GERAN_FREQ_REL,)),
    ("CellRelation", (CELL_REL,)),
])
INFO_SHEETS = ["eNBInfo", "gNBInfo", "LTE_CELL_DATA", "NR_CELL_DATA"]
SECTION_LABEL = {LTE_GPL: "LTE_GPL_AUDIT", NR_GPL: "NR_GPL_AUDIT", FEATURE: "FEATURESTATE", EUTRAN_FREQ: "EUTRANFREQUENCY",
                 SSB_FREQ: "SSBFREQUENCY", EUTRAN_FREQ_REL: "EUTRANFREQRELATION", GUTRAN_FREQ_REL: "GUTRANFREQRELATION",
                 GERAN_FREQ_REL: "GERANFREQRELATION", CELL_REL: "CellRelation", OTHER: "OTHER_MO_AUDIT"}


def _param(it):
    return f"{it.attr}.{it.member}" if it.member else it.attr


def _cell_of(ldn):
    m = re.match(r"^(ENodeBFunction=1,EUtranCell[FT]DD=[^,]+)", ldn)
    return m.group(1) if m else ""


def _band(mo):
    if "_F1_" in mo: return "L2100"
    if "_F3_" in mo: return "L1800"
    if "_F8_" in mo: return "L900"
    if "_T1_" in mo or "_T2_" in mo: return "L23"
    return "Band Error"


EMPTY = "Empty"          # the attribute exists and has no value (empty list, unset reference / string)
NOT_ON_MO = "N/A"        # the attribute does not exist on this MO (other MOs of the table have it)


def _show(v: str) -> str:
    return v if v != "" else EMPTY


def _pipe(pre, post, status):
    if status == "Missing":
        return f"{_show(pre)}|NA"
    return _show(pre) if status == "OK" else f"{_show(pre)}|{_show(post)}"


def _shown_attrs(res, ldn):
    """every configurable PRE attribute of an MO missing in POST, struct members as attr.member, empty ones too"""
    mom = res.post.mom
    key = mom.class_key(ldn)
    cls = classes_of(ldn)[-1]
    out = OrderedDict()
    for a, v in res.pre_mos[ldn].items():
        info = mom.attr(key, a)
        if (info and ({"readOnly", "key"} & info[2])) or a.lower() == (cls[0].lower() + cls[1:] + "Id").lower():
            continue
        if isinstance(v, dict):
            for m, x in v.items():
                out[f"{a}.{m}"] = _show(text(x))
        else:
            out[a] = _show(text(v))
    return out


def _missing_rows_long(res, sections):
    """MOs PRE has and POST has not, one row per configurable PRE attribute (V2: 'Cell is not Found in Post')"""
    rows = []
    for ldn in res.missing_mos:
        if section_of(ldn) in sections:
            rows += [(ldn, a, v) for a, v in _shown_attrs(res, ldn).items()]
    return rows


def build_sheets(res: AuditResult, plan: Plan) -> "OrderedDict[str, list[tuple[str, pd.DataFrame]]]":
    node = res.node
    post_cells = {k: v.get("cellId", "") for k, v in res.post.mos.items() if re.match(r"^ENodeBFunction=1,EUtranCell[FT]DD=[^,]+$", k)}
    sheets = OrderedDict()

    # ---------------------------------------------------------------- Summary (V2: cell mapping + status counts)
    def cells(data, names):
        out = {}
        enb = text((data.mos.get("ENodeBFunction=1") or {}).get("eNBId", "NA")) or "NA"
        for k, v in data.mos.items():
            if re.match(r"^ENodeBFunction=1,EUtranCell[FT]DD=[^,]+$", k):
                out[text(v.get("cellId", ""))] = (data.node, k.split("=")[-1], enb)
        return out
    pc, qc = cells(res.pre, None), cells(res.post, None)
    rows = []
    for cid in sorted(set(pc) | set(qc), key=lambda x: (len(x), x)):
        p, q = pc.get(cid, ("NA", "NA", "NA")), qc.get(cid, ("NA", "NA", "NA"))
        rows.append({"Pre SiteId": p[0], "Pre CellName": p[1], "Pre eNBID": p[2], "cellId": cid,
                     "Post SiteId": q[0], "Post eNBID": q[2], "Post CellName": q[1]})
    summary = pd.DataFrame(rows, columns=["Pre SiteId", "Pre CellName", "Pre eNBID", "cellId", "Post SiteId", "Post eNBID", "Post CellName"])

    # Missing = what the correction script creates (gone after it ran); MOs it cannot create (PRE cell without a cell /
    # radio on the POST site and its whole subtree, node-made, site transport/equipment) are counted apart - they stay
    # after every correction run
    counts = []
    for sec, label in SECTION_LABEL.items():
        its = [i for i in res.items if i.section == sec and not i.info]
        miss = [k for k in res.missing_mos if section_of(k) == sec]
        fixed = sum(1 for k in miss if k not in plan.missing_why)
        counts.append({"Audit Section": label, "Total Parameters": len(its),
                       "OK": sum(i.status == "OK" for i in its), "NOT OK": sum(i.status == "NOT OK" for i in its),
                       "Missing": sum(i.status == "Missing" for i in its) + fixed,
                       "Missing - not scriptable": len(miss) - fixed})
    for sheet in INFO_SHEETS:
        its = [i for i in res.items if i.info == sheet]
        counts.append({"Audit Section": f"{sheet} (audit only)", "Total Parameters": len(its),
                       "OK": sum(i.status == "OK" for i in its), "NOT OK": sum(i.status == "NOT OK" for i in its),
                       "Missing": sum(i.status == "Missing" for i in its), "Missing - not scriptable": 0})
    pre_only_cells = [k.split(",")[-1] for k in res.missing_mos if CELL_CLASSES.match(k)]
    cmd = count_commands(plan)
    info = [("PRE node (reference)", res.pre.node), ("POST node (corrected)", res.post.node),
            ("PRE release", res.pre.release or "NA"), ("POST release", res.post.release or "NA"),
            ("MOs compared", f"PRE {len(res.pre.mos)} / POST {len(res.post.mos)}"),
            ("MOs / profiles missing in POST", len(res.missing_mos)), ("MOs only in POST", len(res.new_mos)),
            ("Cross-site name mapping (PRE -> POST)", ", ".join(f"{a} -> {b}" for a, b in res.mapping) or "same site"),
            ("Neighbour MOs of the other site (not compared)", res.skipped_neighbours),
            ("Relation values the node rewrites itself (lastModification LM / MLB / MRO) - differences not compared",
             res.node_managed),
            ("POST cells without a PRE cell - checked against the values common to the PRE cells",
             "; ".join(f"{c.rsplit(',', 1)[-1]} <- {', '.join(t.rsplit('=', 1)[1] for t in tpl)} ({n} MOs, {d} values "
                       "differ between those PRE cells - not checked)" for c, tpl, n, d in res.template_cells) or "none"),
            ("PRE cells without a cell on the POST site (with their subtree: Missing - not scriptable)",
             (f"{len(pre_only_cells)}: " + ", ".join(pre_only_cells) +
              " - cells come from the integration scripts (RF data / radio); use a PRE site with the same cell "
              "configuration to audit them") if pre_only_cells else "none")]
    info += [(f"Correction commands: {s}", n) for s, n in cmd.items()]
    sheets["Summary"] = [("", summary), ("Parameter Status Summary", pd.DataFrame(counts)),
                         ("Audit Information", pd.DataFrame(info, columns=["Item", "Value"]))]

    # ---------------------------------------------------------------- audit-only + wide sheets
    def wide(items, missing_ldns, with_cell):
        # one table per class AND radio function: EUtranFreqRelation / EUtranCellRelation / ... exist in the LTE
        # (ENodeBFunction) and the NR (GNBCUCPFunction) model with different attributes - one table for both left
        # half the columns blank on every row
        def group(ldn):
            c = classes_of(ldn)
            return c[-1], c[0]
        by_cls = OrderedDict()
        for it in items:
            by_cls.setdefault(group(it.ldn), OrderedDict()).setdefault(it.ldn, []).append(it)
        for ldn in missing_ldns:
            by_cls.setdefault(group(ldn), OrderedDict()).setdefault(ldn, [])
        names = {}
        for cls, top in by_cls:
            names.setdefault(cls, []).append(top)
        tables = []
        for (cls, top), mos in by_cls.items():
            rows = []
            for ldn, its in mos.items():
                row = OrderedDict(Node_ID=node, MO=ldn)
                if with_cell:
                    # MOs outside an LTE cell (NR side, neighbour definitions) have no cellId
                    row["cellId"] = text(post_cells.get(_cell_of(ldn), "")) or NOT_ON_MO
                if ldn in res.missing_mos:
                    row["Status"] = "Missing"
                    row.update(_shown_attrs(res, ldn))
                else:
                    bad = [i for i in its if i.status != "OK"]
                    row["Status"] = "NOT OK" if any(i.status == "NOT OK" for i in bad) else ("Missing" if bad else "OK")
                    for i in its:
                        row[_param(i)] = _pipe(i.pre, i.post, i.status)
                rows.append(row)
            df = pd.DataFrame(rows).fillna(NOT_ON_MO)
            tables.append((f"{cls} ({top})" if len(names[cls]) > 1 else cls, df))
        return tables

    for sheet in INFO_SHEETS:
        its = [i for i in res.items if i.info == sheet]
        tables = wide(its, [], with_cell=False)
        for _, df in tables:                       # V2 order: Node_ID, MO, values..., Status
            df.insert(len(df.columns) - 1, "Status", df.pop("Status"))
        if tables:
            sheets[sheet] = tables

    # LTE GPL (V2 columns)
    rows = []
    for it in (i for i in res.items if i.section == LTE_GPL and not i.info):
        rows.append({"Node_ID": node, "MO": it.ldn, "cellId": text(post_cells.get(_cell_of(it.ldn), 0)) or 0,
                     "Parameter": _param(it), "Pre-existing Value": _show(it.pre),
                     "Current value": _show(it.post) if it.status != "Missing" else "NA", "_merge": "both",
                     "Parameter Setting Status": it.status, "Band": _band(it.ldn)})
    for ldn, a, v in _missing_rows_long(res, {LTE_GPL}):
        rows.append({"Node_ID": "Cell is not Found in Post" if _cell_of(ldn) == ldn else node, "MO": ldn, "cellId": 0,
                     "Parameter": a, "Pre-existing Value": v, "Current value": "NA", "_merge": "left_only",
                     "Parameter Setting Status": "Missing", "Band": _band(ldn)})
    if rows:
        sheets["LTE_GPL_AUDIT"] = [("", pd.DataFrame(rows))]

    # NR GPL (V2 columns, V2 spelling)
    rows = [{"Node_ID": node, "MO": it.ldn, "Perameter": _param(it), "pre_value": _show(it.pre),
             "current_value": _show(it.post) if it.status != "Missing" else "NA", "status": it.status}
            for it in res.items if it.section == NR_GPL and not it.info]
    rows += [{"Node_ID": node, "MO": ldn, "Perameter": a, "pre_value": v, "current_value": "NA", "status": "MISSING IN POST"}
             for ldn, a, v in _missing_rows_long(res, {NR_GPL})]
    if rows:
        sheets["NR_GPL_AUDIT"] = [("", pd.DataFrame(rows))]

    # FeatureState (V2 columns)
    rows = []
    feats = [k for k in list(res.pre_mos) + [k for k in res.post.mos if k not in res.pre_mos] if section_of(k) == FEATURE]
    for k in feats:
        p, q = res.pre_mos.get(k), res.post.mos.get(k)
        if p is None:
            continue
        def st(d, a):
            v = text((d or {}).get(a, ""))
            m = re.match(r"^-?\d+\s*\((\w+)\)$", v)
            return m.group(1) if m else (v or "nan")
        row = {"Node_ID": node, "CXC ID": k.split("=")[-1], "description": _show(text(p.get("description", ""))),
               "Pre Existing FeatureState": st(p, "featureState"), "Pre Existing LicenseState": st(p, "licenseState"),
               "Current FeatureState": st(q, "featureState") if q else "nan",
               "Current LicenseState": st(q, "licenseState") if q else "nan"}
        row["Feature setting Status"] = ("Missing" if q is None else
                                         "OK" if (row["Pre Existing FeatureState"], row["Pre Existing LicenseState"]) ==
                                                 (row["Current FeatureState"], row["Current LicenseState"]) else "NOT OK")
        rows.append(row)
    if rows:
        sheets["FEATURESTATE"] = [("", pd.DataFrame(rows))]

    for sheet, secs in WIDE_SHEETS.items():
        its = [i for i in res.items if i.section in secs and not i.info]
        miss = [k for k in res.missing_mos if section_of(k) in secs]
        tables = wide(its, miss, with_cell=sheet != "FREQ_AUDIT")
        if tables:
            sheets[sheet] = tables

    # MOs / profiles missing in POST (+ only in POST)
    rows = []
    for k in res.missing_mos:
        why = plan.missing_why.get(k, "")
        rows.append({"Node_ID": node, "MO": k, "Class": classes_of(k)[-1], "Section": section_of(k),
                     "Status": "Missing in Post", "Creation Script": why or "Scripted (Missing Parameter/Profile Creation)"})
    for k in res.new_mos:
        rows.append({"Node_ID": node, "MO": k, "Class": classes_of(k)[-1], "Section": section_of(k),
                     "Status": "Only in Post", "Creation Script": "-"})
    if rows:
        sheets["MISSING_MO_PROFILE"] = [("", pd.DataFrame(rows))]

    rows = [{"Node_ID": node, "MO": it.ldn, "Parameter": _param(it), "Pre-existing Value": _show(it.pre),
             "Current value": _show(it.post) if it.status != "Missing" else "NA", "Status": it.status}
            for it in res.items if it.section == OTHER and not it.info]
    if rows:
        sheets["OTHER_MO_AUDIT"] = [("", pd.DataFrame(rows))]

    # what the correction scripts leave out and why (no longer written as comments into the scripts)
    rows = [{"Node_ID": node, "Correction Section": sec, "Not scripted": note} for sec, notes in plan.notes.items()
            for note in notes]
    if rows:
        sheets["NOT_SCRIPTED"] = [("", pd.DataFrame(rows))]
    return sheets


def _write_tables(writer, sheet, tables):
    """titled tables one under the other (V2 styling; every table gets its banner)"""
    if len(tables) == 1 and not tables[0][0]:
        df = tables[0][1]
        df.to_excel(writer, sheet_name=sheet, index=False)
        _format_sheet(writer.sheets[sheet], df)
        return
    row, widths = 1, {}
    max_cols = max(len(df.columns) for _, df in tables)
    for n, (title, df) in enumerate(tables):
        if title:
            if sheet not in writer.sheets:
                pd.DataFrame().to_excel(writer, sheet_name=sheet, index=False)
            ws = writer.sheets[sheet]
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=max(len(df.columns), 1))
            cell = ws.cell(row=row, column=1, value=title)
            cell.fill, cell.font, cell.alignment = _TITLE_FILL, _TITLE_FONT, _TITLE_ALIGN
            ws.row_dimensions[row].height = 24
            row += 1
        df.to_excel(writer, sheet_name=sheet, index=False, startrow=row - 1)
        ws = writer.sheets[sheet]
        for c, w in enumerate(_style_block(ws, df, start_row=row, freeze=(n == 0 and not title)), start=1):
            widths[c] = max(widths.get(c, 10), w)
        row += len(df) + 1 + _TABLE_GAP_ROWS
    for c in range(1, max_cols + 1):
        ws.column_dimensions[get_column_letter(c)].width = min(widths.get(c, 10) + 3, 55)


def write_report(sheets, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for sheet, tables in sheets.items():
            tables = [(t, df) for t, df in tables if not df.empty]
            if tables:
                _write_tables(writer, sheet, tables)
